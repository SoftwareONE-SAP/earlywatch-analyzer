"""Whole-workbook checks for the polished minimal look: navy frame only, one font, plain view."""
from __future__ import annotations

import re
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText

from ewa_pipeline.report import excel_style as xs
from ewa_pipeline.report.excel_generator import generate
from ewa_pipeline.report.schemas import (
    AnalysisTaskRecord, CrossReference, ReconciliationRow, SectionCoverage, SystemProfile,
)
from test_excel_generator import build_result


def full_result():
    """A result that makes every tab appear."""
    result = build_result()
    result.executive_summary = "Two things need attention."
    result.top_5_priority_actions = ["Tune the statement (F-002) and patch the kernel (F-001)."]
    result.system_profile = SystemProfile(sid="PRD", product="SAP S/4HANA", product_version="2022",
                                          period_start="01.05.2026", period_end="31.05.2026")
    result.cross_references = [CrossReference(
        title="Kernel and SQL", related_findings=["F-001", "F-002"],
        correlation_description="Both slow the system.", combined_impact="Slow.",
        recommended_action="Fix both.")]
    result.coverage = [
        SectionCoverage(section_id="1", section_title="Software Configuration", analysis_task_id="T1",
                        status="analyzed"),
        SectionCoverage(section_id="4", section_title="Database", analysis_task_id="T2", status="analyzed"),
    ]
    result.analysis_tasks = [AnalysisTaskRecord(
        task_id="T1", title="Task", source_section_ids=["1"], source_content_chars=10, origin="plan")]
    result.reconciliation = [
        ReconciliationRow(source="Alert", verdict="Agree", sap_topic="Performance", sap_item="Slow SQL",
                          sap_rating="RED", our_section_id="2", our_section_title="Performance Indicators",
                          our_rating="RED", finding_ids=["F-002"]),
        ReconciliationRow(source="Check", verdict="Not covered", sap_topic="Database", sap_rating="YELLOW",
                          our_section_id="4"),
    ]
    return result


TREE = {"structure": [{"title": "Overview", "start_index": 1, "end_index": 3, "summary": "s",
                       "nodes": [{"title": "Child", "start_index": 2, "end_index": 3, "summary": "c"}]}]}


class WholeWorkbookStyleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            generate(full_result(), path, tree=TREE)
            cls.wb = load_workbook(path, rich_text=True)

    def test_every_expected_tab_is_present(self):
        names = self.wb.sheetnames
        for name in ("Priority Actions", "Findings", "SAP vs SoftwareONE",
                     "Risk Heatmap", "Document Structure",
                     "Token Usage", "Analysis Tasks", "Source Coverage"):
            self.assertIn(name, names)

    def _navy_rows(self, ws):
        """Title-band rows and table-header rows: the only rows allowed to have a fill."""
        rows = set()
        if ws.title in ("Priority Actions", "SAP vs SoftwareONE", "Risk Heatmap"):
            rows.add(1)
        if ws.freeze_panes:
            rows.add(ws[ws.freeze_panes].row - 1)  # main table header
        for row in ws.iter_rows():
            values = [c.value for c in row if c.value is not None]
            if "Alerts" in values:  # SAP vs Us summary table
                rows.add(row[0].row)
        return rows

    def test_fills_only_navy_on_title_bands_and_header_rows(self):
        for ws in self.wb.worksheets:
            navy_rows = self._navy_rows(ws)
            for cell in ws._cells.values():
                where = f"{ws.title}!{cell.coordinate}"
                if cell.fill.fill_type in (None, ""):
                    continue
                self.assertEqual(cell.fill.fill_type, "solid", where)
                if ws.title == "Risk Heatmap" and cell.row not in navy_rows:
                    # deliberate heatmap exception: grid cells are filled with the severity colours
                    self.assertIn(cell.fill.start_color.rgb[-6:], xs.SEVERITY_COLOURS.values(), where)
                    continue
                self.assertEqual(cell.fill.start_color.rgb[-6:], "1F3864", where)
                self.assertIn(cell.row, navy_rows, where)

    def test_every_navy_header_cell_has_white_text(self):
        for ws in self.wb.worksheets:
            for cell in ws._cells.values():
                if cell.fill.fill_type == "solid" and cell.value is not None:
                    self.assertEqual(cell.font.color.rgb[-6:], "FFFFFF", f"{ws.title}!{cell.coordinate}")

    def test_body_cells_have_no_fill(self):
        for ws in self.wb.worksheets:
            navy_rows = self._navy_rows(ws)
            for cell in ws._cells.values():
                if cell.row in navy_rows:
                    continue
                if ws.title == "Risk Heatmap" and cell.fill.fill_type == "solid":
                    # deliberate heatmap exception: only the severity colours may fill a body cell
                    self.assertIn(cell.fill.start_color.rgb[-6:], xs.SEVERITY_COLOURS.values(),
                                  f"{ws.title}!{cell.coordinate}")
                    continue
                self.assertIn(cell.fill.fill_type, (None, ""), f"{ws.title}!{cell.coordinate}")

    def test_every_font_is_aptos_including_rich_text_runs(self):
        for ws in self.wb.worksheets:
            for cell in ws._cells.values():
                self.assertEqual(cell.font.name, "Aptos", f"{ws.title}!{cell.coordinate}")
                if isinstance(cell.value, CellRichText):
                    for part in cell.value:
                        if not isinstance(part, str):
                            self.assertEqual(part.font.rFont, "Aptos", f"{ws.title}!{cell.coordinate}")
                        else:
                            self.fail(f"plain run inside rich text at {ws.title}!{cell.coordinate}")

    def test_zoom_100_and_no_gridlines_on_every_tab(self):
        for ws in self.wb.worksheets:
            self.assertEqual(ws.sheet_view.zoomScale, 100, ws.title)
            self.assertFalse(ws.sheet_view.showGridLines, ws.title)

    def test_tables_repeat_their_header_row_when_printed(self):
        for ws in self.wb.worksheets:
            if ws.freeze_panes:
                self.assertTrue(ws.print_title_rows, ws.title)

    def test_all_tabs_print_a4_landscape_one_page_wide(self):
        for ws in self.wb.worksheets:
            self.assertEqual(ws.page_setup.orientation, "landscape", ws.title)
            self.assertEqual(int(ws.page_setup.paperSize), 9, ws.title)
            self.assertEqual(ws.page_setup.fitToWidth, 1, ws.title)

    def test_every_internal_link_lands_on_an_existing_sheet_and_row(self):
        finding_ids = {f.id for da in full_result().domain_analyses for f in da.findings}
        for ws in self.wb.worksheets:
            for cell in ws._cells.values():
                link = cell.hyperlink
                target = (link.target or link.location or "") if link else ""
                if not target.startswith("#"):
                    continue
                where = f"{ws.title}!{cell.coordinate}"
                m = re.fullmatch(r"#'(.+)'!([A-Z]+)(\d+)", target)
                self.assertIsNotNone(m, f"{where}: {target}")
                self.assertIn(m.group(1), self.wb.sheetnames, where)
                self.assertLessEqual(int(m.group(3)), self.wb[m.group(1)].max_row, where)
                if m.group(1) == "Findings":  # a link into Findings lands on a finding's row
                    self.assertIn(self.wb["Findings"][f"A{m.group(3)}"].value, finding_ids, where)


if __name__ == "__main__":
    unittest.main()
