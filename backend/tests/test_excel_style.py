"""Minimal-report Excel style helpers."""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook

from ewa_pipeline.report import excel_style as xs


class ExcelStyleTests(unittest.TestCase):
    def test_default_font_is_font_name(self):
        wb = xs.init_workbook()
        self.assertEqual(wb._fonts[0].name, xs.FONT_NAME)
        ws = wb.active
        ws["A1"] = "x"
        self.assertEqual(ws["A1"].font.name, xs.FONT_NAME)
        self.assertEqual(ws["B5"].font.name, xs.FONT_NAME)  # empty cell too

    def test_font_sets_name_and_size(self):
        f = xs.font(bold=True)
        self.assertEqual((f.name, f.sz, f.b), (xs.FONT_NAME, 11, True))

    def test_header_row_is_navy_with_white_text_and_thin_navy_border(self):
        ws = xs.init_workbook().active
        xs.header_row(ws, 1, ["Name", "Count"], numeric_cols=(2,))
        a, b = ws["A1"], ws["B1"]
        self.assertTrue(a.font.b)
        self.assertEqual(a.font.color.rgb[-6:], "FFFFFF")
        self.assertEqual(a.fill.fill_type, "solid")
        self.assertEqual(a.fill.start_color.rgb[-6:], "1F3864")
        self.assertEqual(a.border.bottom.style, "thin")
        self.assertEqual(a.border.bottom.color.rgb[-6:], "1F3864")
        self.assertEqual(a.alignment.horizontal, "left")
        self.assertEqual(b.alignment.horizontal, "right")

    def test_table_style_borders_and_filter(self):
        ws = xs.init_workbook().active
        xs.header_row(ws, 1, ["A", "B"])
        ws.append(["x", 1])
        xs.apply_table_style(ws, last_row=2, last_col=2)
        self.assertEqual(ws["A2"].border.bottom.style, "thin")
        self.assertEqual(ws["A2"].alignment.vertical, "top")
        self.assertEqual(ws.auto_filter.ref, "A1:B2")

    def test_sheet_view_and_print(self):
        ws = xs.init_workbook().active
        xs.apply_sheet_view(ws, freeze="A2")
        xs.apply_print_setup(ws, repeat_row=1)
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertEqual(ws.sheet_view.zoomScale, 100)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(ws.page_setup.paperSize, 9)  # 9 = A4
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertEqual(ws.page_setup.fitToHeight, 0)
        self.assertEqual(ws.print_title_rows, "$1:$1")

    def test_text_helpers_read_back_as_plain_text(self):
        self.assertEqual(str(xs.rating_text("RED")), f"{xs.DOT} Red")
        self.assertEqual(str(xs.rating_text("yellow")), f"{xs.DOT} Yellow")
        self.assertEqual(str(xs.severity_text("High")), f"{xs.DOT} High")
        self.assertEqual(str(xs.sure_dots("High")), "●●● High")
        self.assertEqual(str(xs.sure_dots("Medium")), "●●○ Medium")
        self.assertEqual(str(xs.sure_dots("Low")), "●○○ Low")
        self.assertEqual(str(xs.sure_dots("")), xs.DASH)
        self.assertEqual(str(xs.sure_dots(None)), xs.DASH)
        self.assertEqual(str(xs.rating_text("")), xs.DASH)
        self.assertEqual(str(xs.muted("n/a")), "n/a")

    def test_rich_text_survives_save_and_reads_plain(self):
        wb = xs.init_workbook()
        ws = wb.active
        ws["A1"] = xs.severity_text("High")
        ws["A2"] = xs.sure_dots("Medium")
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "t.xlsx"
            wb.save(path)
            plain = load_workbook(path).active
            self.assertEqual(plain["A1"].value, f"{xs.DOT} High")
            self.assertEqual(plain["A2"].value, "●●○ Medium")
            if xs.RICH:
                rich = load_workbook(path, rich_text=True).active
                self.assertEqual(str(rich["A1"].value), f"{xs.DOT} High")


if __name__ == "__main__":
    unittest.main()
