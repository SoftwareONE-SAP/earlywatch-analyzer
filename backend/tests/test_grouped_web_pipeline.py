"""Synthetic grouped analysis through the actual HTML, graph and Excel workflow."""
from __future__ import annotations

import json
import re
import sys
import unittest
from contextlib import ExitStack
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from ewa_pipeline.agents import orchestrator
from ewa_pipeline.agents.runner import AnalysisFailedError
from ewa_pipeline.report.schemas import (
    CrossReferenceList, DomainAnalysis, Finding, GroupedDomainAnalysis, OrchestratorPlan,
    Remediation, SynthesisResult,
)
from ewa_pipeline.agents.prompts import DOMAIN_ANALYST_PROMPT, GROUPED_DOMAIN_ANALYST_PROMPT
from ewa_pipeline.services.pipeline import run_pipeline
from test_cli_analysis_settings import configured


HTML = """<article><h1>Database</h1><h2>SQL A</h2>
<p>SQL A runtime 95 ms; statement hash A123.</p><h2>SQL B</h2>
<p>SQL B runtime 35 ms; statement hash B456.</p></article>"""


class GroupedWebPipelineTests(unittest.TestCase):
    @staticmethod
    def rich_finding():
        return Finding(
            id="F001", title="Slow statement", severity="Medium",
            description="Likely missing index", evidence="[YELLOW] runtime 95 ms",
            impact="Slower dialog response",
            remediation=Remediation(action="Check the plan", sap_transactions=["ST05"],
                                    effort_estimate="Low", priority="Short-term"),
            confidence="High", confidence_reason="CONF_REASON_MARKER",
            to_confirm="TO_CONFIRM_MARKER", risk_types=["Performance"],
            subject="SUBJECT_MARKER_A123", due_date="2026-12-31",
        )

    def run_synthetic(self, directory, *, mode="complete", config=None, planner_fails=False):
        config = config or configured()
        source = Path(directory) / "synthetic.html"
        source.write_text(HTML, encoding="utf-8")
        output = Path(directory) / "synthetic.xlsx"
        calls = []
        self.last_calls = calls

        def invoke(schema, messages):
            prompt = "\n".join(str(message.content) for message in messages)
            calls.append((schema, prompt))
            if schema is OrchestratorPlan:
                if planner_fails:
                    raise RuntimeError("Synthetic planner failure")
                parsed = OrchestratorPlan(tasks=[{"section_ids": ["0002", "0003"],
                    "section_title": "Database SQL", "analysis_focus": "Compare related statements"}])
            elif schema in (DomainAnalysis, GroupedDomainAnalysis):
                ids = re.findall(r"Source Section ID: (\d+)", prompt)
                if not ids:
                    ids = re.findall(r"^ID: (\d+)", prompt, re.MULTILINE)
                self.assertTrue(ids)
                if mode == "exception":
                    raise RuntimeError("synthetic provider payload must never escape")
                if mode == "missing":
                    ids = ids[:1]
                elif mode == "duplicate":
                    ids = [ids[0], ids[0]]
                elif mode == "unknown":
                    ids = ["9999"]
                analyses = [DomainAnalysis(section_id=sid, section_title="Untrusted model title",
                            findings=[self.rich_finding()] if mode == "findings" else [],
                            overall_health="Warning" if mode == "findings" else "Healthy") for sid in ids]
                parsed = GroupedDomainAnalysis(analyses=analyses) if schema is GroupedDomainAnalysis else analyses[0]
            elif schema is CrossReferenceList:
                parsed = CrossReferenceList(items=[])
            elif schema is SynthesisResult:
                if mode == "synthesis_failure":
                    raise RuntimeError("synthetic synthesis failure")
                parsed = SynthesisResult(executive_summary="Synthetic grouped review.",
                                         overall_system_health="Healthy", top_5_priority_actions=[])
            else:
                raise AssertionError(schema)
            return {"parsed": parsed,
                    "raw": SimpleNamespace(usage_metadata={"input_tokens": 100, "output_tokens": 20,
                        "input_token_details": {"cache_read": 10, "cache_write_tokens": 5}})}

        def structured(schema, **kwargs):
            chain = Mock()
            chain.invoke.side_effect = lambda messages: invoke(schema, messages)
            return chain

        with ExitStack() as stack:
            for name in ("get_orchestrator_model", "get_subagent_model", "get_cross_ref_model", "get_synthesis_model"):
                factory = stack.enter_context(patch.object(orchestrator, name))
                factory.return_value.with_structured_output.side_effect = structured
            result, tracker, artifacts = run_pipeline(config=config, input_path=source, output_path=output)
        return result, tracker, artifacts, calls

    def test_shared_request_keeps_sources_usage_and_workbook_attribution(self):
        with TemporaryDirectory() as directory:
            result, tracker, artifacts, calls = self.run_synthetic(directory)
            self.assertFalse(result.degraded)
            self.assertEqual([analysis.section_id for analysis in result.domain_analyses], ["0002", "0003"])
            self.assertEqual([analysis.section_title for analysis in result.domain_analyses], ["SQL A", "SQL B"])
            self.assertEqual([source.status for source in result.coverage], ["analyzed", "analyzed"])
            self.assertEqual(len(result.analysis_tasks), 1)
            self.assertEqual(result.analysis_tasks[0].source_section_ids, ["0002", "0003"])
            self.assertEqual(result.analysis_tasks[0].context_section_ids, ["0001"])
            grouped = [prompt for schema, prompt in calls if schema is GroupedDomainAnalysis]
            self.assertEqual(len(grouped), 1)
            self.assertIn("A123", grouped[0])
            self.assertIn("B456", grouped[0])
            self.assertEqual(tracker._entries["phase1_domain_analysis"].calls, 1)
            self.assertEqual(tracker._entries["phase1_domain_analysis"].input_tokens, 100)
            self.assertEqual(result.token_usage.phase1_cache_write_tokens, 5)
            persisted = json.loads(artifacts.result_path.read_text(encoding="utf-8"))
            self.assertEqual(len(persisted["coverage"]), 2)
            workbook = load_workbook(artifacts.output_path)
            self.assertIn("Analysis Tasks", workbook.sheetnames)
            self.assertIn("Source Coverage", workbook.sheetnames)
            self.assertEqual(workbook["Analysis Tasks"]["C2"].value, "0002, 0003")
            self.assertEqual(workbook["Source Coverage"]["A3"].value, "0003")
            workbook.close()

    def test_portal_count_limit_splits_shared_task(self):
        config = configured().model_copy(update={"max_task_sections": 1, "max_parent_context_chars": 0})
        with TemporaryDirectory() as directory:
            result, tracker, _, calls = self.run_synthetic(directory, config=config)
        self.assertEqual(len(result.analysis_tasks), 2)
        self.assertTrue(all(task.context_content_chars == 0 for task in result.analysis_tasks))
        self.assertEqual(tracker._entries["phase1_domain_analysis"].calls, 2)
        self.assertFalse(any(schema is GroupedDomainAnalysis for schema, _ in calls))

    def test_planner_failure_falls_back_to_related_hierarchy_group(self):
        with TemporaryDirectory() as directory:
            result, _, _, _ = self.run_synthetic(directory, planner_fails=True)
        self.assertEqual(len(result.analysis_tasks), 1)
        self.assertEqual(result.analysis_tasks[0].origin, "default")
        self.assertFalse(result.degraded)

    def test_missing_source_marks_partial_workbook_unknown(self):
        with TemporaryDirectory() as directory:
            result, _, artifacts, _ = self.run_synthetic(directory, mode="missing")
            self.assertTrue(result.degraded)
            self.assertEqual(result.overall_system_health, "Unknown")
            self.assertEqual([source.status for source in result.coverage], ["analyzed", "failed"])
            workbook = load_workbook(artifacts.output_path)
            self.assertIn("UNKNOWN", workbook["Executive Summary"]["B2"].value)
            self.assertEqual(workbook["Source Coverage"]["D3"].value, "failed")
            workbook.close()

    def test_duplicate_unknown_or_failed_group_cannot_publish_healthy_result(self):
        for mode in ("duplicate", "unknown", "exception"):
            with self.subTest(mode=mode), TemporaryDirectory() as directory:
                with self.assertRaises(AnalysisFailedError):
                    self.run_synthetic(directory, mode=mode)
                self.assertFalse((Path(directory) / "synthetic.xlsx").exists())

    def test_oversized_sources_never_invoke_specialists_or_publish_workbook(self):
        config = configured().model_copy(update={"max_task_content_chars": 1})
        with TemporaryDirectory() as directory:
            with self.assertRaises(AnalysisFailedError):
                self.run_synthetic(directory, config=config)
            self.assertFalse((Path(directory) / "synthetic.xlsx").exists())
        self.assertEqual([schema for schema, _ in self.last_calls], [OrchestratorPlan])

    def test_synthesis_failure_is_unknown_with_preserved_source_results(self):
        with TemporaryDirectory() as directory:
            result, _, _, _ = self.run_synthetic(directory, mode="synthesis_failure")
        self.assertTrue(result.degraded)
        self.assertEqual(result.overall_system_health, "Unknown")
        self.assertEqual(len(result.domain_analyses), 2)

    def test_analyst_prompts_format_and_carry_finding_field_guidance(self):
        single = DOMAIN_ANALYST_PROMPT.format(
            section_title="T", section_id="0001", analysis_focus="F", skill_context="S", content="C")
        grouped = GROUPED_DOMAIN_ANALYST_PROMPT.format(
            section_title="T", section_ids="0001", analysis_focus="F", skill_context="S", content="C")
        for prompt in (single, grouped):
            for text in ("Security, Availability, Performance, Compliance, Cost, Data integrity, "
                         "Supportability, Operations", "confidence_reason", "to_confirm",
                         "subject", "due_date", "YYYY-MM-DD", "verbatim"):
                self.assertIn(text, prompt)

    def test_later_calls_omit_bulky_finding_fields_but_keep_correlation_fields(self):
        with TemporaryDirectory() as directory:
            result, _, _, calls = self.run_synthetic(directory, mode="findings")
        self.assertTrue(any(f.confidence_reason for a in result.domain_analyses for f in a.findings))
        later = [prompt for schema, prompt in calls if schema in (CrossReferenceList, SynthesisResult)]
        self.assertEqual(len(later), 2)
        for prompt in later:
            for key in ("confidence_reason", "to_confirm", "due_date",
                        "CONF_REASON_MARKER", "TO_CONFIRM_MARKER", "2026-12-31"):
                self.assertNotIn(key, prompt)
            for key in ('"subject"', "SUBJECT_MARKER_A123", '"risk_types"', '"confidence"'):
                self.assertIn(key, prompt)


if __name__ == "__main__":
    unittest.main()
