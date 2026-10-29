"""The Excel workbook has a Findings tab and no per-section tabs."""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ewa_pipeline.report.excel_generator import _finding_rows, _refs_in_text, generate
from ewa_pipeline.report.schemas import (
    AnalysisResult, CrossReference, DomainAnalysis, Finding, ReconciliationRow, Remediation, SectionCoverage,
    SystemProfile,
)
from ewa_pipeline.tracking.token_tracker import TokenUsage


def _remediation(action: str) -> Remediation:
    return Remediation(
        action=action,
        sap_transactions=["SPAM"],
        effort_estimate="Medium",
        priority="Short-term",
    )


def old_style_finding() -> Finding:
    """Shaped like a finding saved before the new fields existed."""
    return Finding(
        id="F-001",
        title="Kernel release out of maintenance",
        severity="Medium",
        description="The kernel is old.",
        evidence="Kernel 753 [RED]",
        impact="No security patches.",
        remediation=_remediation("Upgrade the kernel."),
    )


def new_style_finding() -> Finding:
    return Finding(
        id="F-002",
        title="Expensive SQL statement",
        severity="Critical",
        description="One statement dominates database time.",
        evidence="Statement hash 4f3a9c uses 61% of DB time",
        impact="Slow dialog response.",
        remediation=_remediation("Tune the statement."),
        confidence="High",
        confidence_reason="The report rates it RED.",
        to_confirm="Check ST04 / DBACOCKPIT.",
        risk_types=["Performance", "Cost"],
        subject="4f3a9c",
        due_date="2026-12-31",
    )


def build_result() -> AnalysisResult:
    return AnalysisResult(
        domain_analyses=[
            DomainAnalysis(
                section_title="Software Configuration",
                section_id="1",
                findings=[old_style_finding()],
                overall_health="Warning",
            ),
            DomainAnalysis(
                section_title="Performance Indicators",
                section_id="2",
                findings=[new_style_finding()],
                overall_health="Critical",
            ),
        ],
        cross_references=[],
        executive_summary="Summary.",
        overall_system_health="Critical",
        top_5_priority_actions=["Tune the statement."],
        token_usage=TokenUsage(),
    )


def _headers(ws, row: int) -> list:
    return [c.value for c in ws[row]]


class ExcelGeneratorTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        path = Path(self._tmp.name) / "out.xlsx"
        generate(build_result(), path)
        self.wb = load_workbook(path)

    def test_findings_follows_priority_actions(self):
        names = self.wb.sheetnames
        self.assertEqual(names[0], "Priority Actions")
        self.assertEqual(names[1], "Findings")
        self.assertNotIn("How to read this workbook", names)
        for gone in ("Software Configuration", "Performance Indicators", "Action Plan", "Token Usage"):
            self.assertNotIn(gone, names)

    def test_findings_rows_sorted_by_severity(self):
        ws = self.wb["Findings"]
        headers = _headers(ws, 1)
        self.assertEqual(headers[:9], ["Ref.", "Severity", "Title", "What it means", "Recommended action",
                                       "Priority", "Effort", "How sure we are", "Section"])
        self.assertEqual(len(headers), 19)
        self.assertIn("What the SAP report says", headers)
        self.assertEqual(headers[-2:], ["Reviewer note: why check first", "Reviewer note: search query"])
        self.assertEqual(ws.max_row, 3)  # header + exactly 2 findings

        rows = [dict(zip(headers, [c.value for c in ws[r]])) for r in (2, 3)]
        self.assertEqual([r["Severity"] for r in rows], ["● Critical", "● Medium"])

        new, old = rows
        self.assertEqual(new["Ref."], "F-002")
        self.assertEqual(new["Section"], "Performance Indicators")
        self.assertEqual(new["Risk area"], "Performance, Cost")
        self.assertEqual(new["How sure we are"], "●●● High")
        self.assertEqual(new["Why"], "The report rates it RED.")
        self.assertEqual(new["How to confirm"], "Check ST04 / DBACOCKPIT.")
        self.assertEqual(new["What the SAP report says"], "Statement hash 4f3a9c uses 61% of DB time")
        self.assertEqual(new["What it means"], "One statement dominates database time.")
        self.assertEqual(new["Due date"], "2026-12-31")
        self.assertEqual(new["Check latest SAP guidance first?"], "No")

        # Old-style finding: new fields render as empty cells (or a dash for How sure).
        self.assertIn(old["Risk area"], ("", None))
        self.assertEqual(old["How sure we are"], "—")
        self.assertIn(old["How to confirm"], ("", None))
        self.assertEqual(old["What the SAP report says"], "Kernel 753 [RED]")
        self.assertEqual(old["Recommended action"], "Upgrade the kernel.")
        self.assertEqual(old["Priority"], "Short-term")

    def test_detail_columns_are_grouped_and_collapsed(self):
        ws = self.wb["Findings"]
        for col in range(10, 20):
            dim = ws.column_dimensions[get_column_letter(col)]
            self.assertTrue(dim.hidden, col)
            self.assertEqual(dim.outlineLevel, 1, col)
        for col in range(1, 10):
            dim = ws.column_dimensions[get_column_letter(col)]
            self.assertFalse(dim.hidden, col)
            self.assertEqual(dim.outlineLevel, 0, col)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:S3")
        self.assertIsNone(ws["A2"].hyperlink)  # no link from Ref. to a section tab any more

    def test_reviewer_notes_columns_are_grouped_and_collapsed(self):
        ws = self.wb["Findings"]
        headers = _headers(ws, 1)
        for name in ("Reviewer note: why check first", "Reviewer note: search query"):
            dim = ws.column_dimensions[get_column_letter(headers.index(name) + 1)]
            self.assertTrue(dim.hidden, name)
            self.assertEqual(dim.outlineLevel, 1, name)
        # Severity (column B) is never hidden.
        self.assertFalse(ws.column_dimensions["B"].hidden)

    def test_section_titled_findings_does_not_create_a_second_tab(self):
        result = build_result()
        result.domain_analyses[0].section_title = "Findings"
        path = Path(self._tmp.name) / "collide.xlsx"
        generate(result, path)
        names = load_workbook(path).sheetnames
        self.assertEqual(len(names), len(set(names)))
        self.assertEqual(names.count("Findings"), 1)
        self.assertEqual(names[1], "Findings")

    def test_priority_actions_is_first_and_active(self):
        self.assertEqual(self.wb.active.title, "Priority Actions")
        selected = [ws.title for ws in self.wb.worksheets if ws.sheet_view.tabSelected]
        self.assertEqual(selected, ["Priority Actions"])

    def test_no_tab_is_hidden_and_there_are_no_technical_tabs(self):
        for ws in self.wb.worksheets:
            self.assertEqual(ws.sheet_state, "visible", ws.title)
        for gone in ("Document Structure", "Token Usage", "Analysis Tasks", "Source Coverage"):
            self.assertNotIn(gone, self.wb.sheetnames)


def _finding(fid: str, severity: str) -> Finding:
    return Finding(
        id=fid, title=f"Finding {fid}", severity=severity,
        description="d", evidence="e", impact="i", remediation=_remediation("Fix."),
    )


class FindingRowsTests(unittest.TestCase):
    def test_rows_follow_severity_order_and_start_at_two(self):
        row_by_id, first_by_section = _finding_rows(build_result())
        self.assertEqual(row_by_id, {"F-002": 2, "F-001": 3})  # Critical before Medium
        self.assertEqual(first_by_section, {"2": 2, "1": 3})

    def test_duplicate_id_keeps_first_row(self):
        result = build_result()
        result.domain_analyses[0].findings.append(new_style_finding())  # a second F-002, in section 1
        row_by_id, _ = _finding_rows(result)
        self.assertEqual(row_by_id["F-002"], 2)


class RefsInTextTests(unittest.TestCase):
    def _result(self, ids_by_section):
        result = build_result()
        result.domain_analyses = [
            DomainAnalysis(section_title=f"S{i}", section_id=str(i), overall_health="Warning",
                           findings=[_finding(fid, "High") for fid in ids])
            for i, ids in enumerate(ids_by_section, 1)]
        return result

    def test_plain_ids_in_order_of_appearance(self):
        result = self._result([["0052-F001", "0052-F002"], ["0053-F001"]])
        text = "Fix 0053-F001 first, then 0052-F001."
        self.assertEqual(_refs_in_text(text, result), ["0053-F001", "0052-F001"])

    def test_range_with_en_dash_brings_in_the_findings_between(self):
        result = self._result([["0015-F001", "0015-F002", "0015-F003", "0015-F004"]])
        self.assertEqual(_refs_in_text("See 0015-F001–0015-F003 for details.", result),
                         ["0015-F001", "0015-F002", "0015-F003"])
        self.assertEqual(_refs_in_text("0015-F001-0015-F003", result),
                         ["0015-F001", "0015-F002", "0015-F003"])

    def test_range_across_sections_is_not_filled_in(self):
        result = self._result([["A-F001", "A-F002"], ["B-F001"]])
        self.assertEqual(_refs_in_text("A-F001 – B-F001", result), ["A-F001", "B-F001"])

    def test_unknown_ids_are_skipped_and_partial_words_do_not_match(self):
        result = self._result([["0052-F001"]])
        self.assertEqual(_refs_in_text("0099-F001 and 0052-F001 and 0052-F0011", result), ["0052-F001"])

    def test_f_dash_style_ids(self):
        self.assertEqual(_refs_in_text("Tune F-002, then upgrade (F-001).", build_result()), ["F-002", "F-001"])

    def test_no_repeats(self):
        self.assertEqual(_refs_in_text("F-002, F-001 and F-002 again", build_result()), ["F-002", "F-001"])

    def test_no_findings_gives_nothing(self):
        result = build_result()
        for da in result.domain_analyses:
            da.findings = []
        self.assertEqual(_refs_in_text("F-001", result), [])


def _row(source, verdict, **kw) -> ReconciliationRow:
    return ReconciliationRow(source=source, verdict=verdict, **kw)


class SapVsUsTests(unittest.TestCase):
    def _generate(self, result):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "out.xlsx"
        generate(result, path)
        return load_workbook(path)

    def test_sap_vs_us_sheet_position_rows_and_counts(self):
        result = build_result()
        result.reconciliation = [
            _row("Alert", "Agree", sap_topic="Alerts", sap_item="Kernel old", sap_rating="RED",
                 our_section_id="1", our_section_title="Software Configuration",
                 our_rating="RED", finding_ids=["F-001", "F-003"], match_score=0.97, match_method="Decision-1"),
            _row("Alert", "Not covered", sap_item="Odd alert", sap_rating="YELLOW"),
            _row("Check", "Agree", sap_topic="Performance", sap_rating="GREEN", our_rating="GREEN"),
            _row("Check", "We rated higher", sap_topic="Hardware", sap_rating="YELLOW", our_rating="RED"),
        ]
        wb = self._generate(result)
        self.assertEqual(wb.sheetnames[:3], ["Priority Actions", "Findings", "SAP vs SoftwareONE"])
        ws = wb["SAP vs SoftwareONE"]

        counts = {ws.cell(r, 1).value: (ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 6).value)
                  for r in range(5, 9)}
        self.assertEqual(counts["Agree"], (1, 1, 2))
        self.assertEqual(counts["We rate it more serious"], (0, 1, 1))
        self.assertEqual(counts["Not found in our analysis"], (1, 0, 1))

        hdr_row = next(r for r in range(1, 15) if ws.cell(r, 1).value == "SAP list")
        headers = _headers(ws, hdr_row)
        self.assertEqual(headers[:5], ["SAP list", "SAP topic", "What SAP says", "SAP rating", "SoftwareONE rating"])
        self.assertEqual(ws.max_row - hdr_row, 4)  # one row per reconciliation row
        first = dict(zip(headers, [c.value for c in ws[hdr_row + 1]]))
        self.assertEqual(first["Result"], "Agree")
        self.assertEqual(first["SAP rating"], "● Red")
        self.assertEqual(first["SoftwareONE rating"], "● Red")
        self.assertEqual(first["How sure we are"], "●●● High")
        self.assertEqual(first["Our finding"], "Kernel release out of maintenance (+1 more)")
        target = ws.cell(hdr_row + 1, headers.index("Our finding") + 1).hyperlink
        self.assertEqual(target.target or target.location, "#'Findings'!A3")  # Medium, after Critical F-002
        self.assertIsNone(ws.cell(hdr_row + 1, headers.index("Report section") + 1).hyperlink)
        self.assertEqual(first["Reviewer note: refs"], "F-001, F-003")
        self.assertEqual(first["Reviewer note: score"], 0.97)
        self.assertEqual(first["Reviewer note: matched by"], "Decision-1")
        self.assertEqual(ws.freeze_panes, f"A{hdr_row + 1}")

        # Bold only for the results a reader should look at; reviewer notes are collapsed.
        result_col = headers.index("Result") + 1
        self.assertFalse(ws.cell(hdr_row + 1, result_col).font.b)  # Agree
        self.assertTrue(ws.cell(hdr_row + 4, result_col).font.b)   # more serious
        for name in headers:
            if name.startswith("Reviewer note"):
                self.assertTrue(ws.column_dimensions[get_column_letter(headers.index(name) + 1)].hidden, name)

        # Fills only on the title band and the two header rows; no matching jargon in anything visible.
        for row in ws.iter_rows():
            for cell in row:
                if cell.row in (1, hdr_row) or (cell.row == 4 and cell.column <= 6):
                    self.assertEqual(cell.fill.fill_type, "solid", cell.coordinate)
                    self.assertTrue(str(cell.fill.start_color.rgb).endswith("1F3864"), cell.coordinate)
                else:
                    self.assertIn(cell.fill.fill_type, (None, ""), cell.coordinate)
                if cell.value is not None and not ws.column_dimensions[cell.column_letter].hidden:
                    for jargon in ("Decision", "p=", "overlap"):
                        self.assertNotIn(jargon, str(cell.value), cell.coordinate)

    def test_both_sheets_ordered_when_both_present(self):
        result = build_result()
        result.coverage = [SectionCoverage(section_id="9", section_title="Clean", status="analyzed")]
        result.reconciliation = [_row("Check", "Agree", sap_rating="GREEN")]
        names = self._generate(result).sheetnames
        self.assertEqual(names[1:3], ["Findings", "SAP vs SoftwareONE"])
        self.assertNotIn("No Issues Found", names)

    def test_old_payload_has_no_sap_sheet(self):
        wb = self._generate(build_result())
        self.assertNotIn("SAP vs SoftwareONE", wb.sheetnames)
        self.assertNotIn("No Issues Found", wb.sheetnames)


class MinimalReportBaseStyleTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "out.xlsx"
        generate(build_result(), path)
        self.wb = load_workbook(path)

    def test_every_tab_is_plain_100_percent_no_gridlines(self):
        for ws in self.wb.worksheets:
            self.assertEqual(ws.sheet_view.zoomScale, 100, ws.title)
            self.assertFalse(ws.sheet_view.showGridLines, ws.title)

    def test_table_headers_are_navy_with_bold_white_text(self):
        ws = self.wb["Findings"]
        for cell in ws[1]:
            self.assertTrue(cell.font.b)
            self.assertEqual(cell.font.name, "Aptos")
            self.assertEqual(cell.font.color.rgb[-6:], "FFFFFF")
            self.assertEqual(cell.fill.fill_type, "solid")
            self.assertEqual(cell.fill.start_color.rgb[-6:], "1F3864")
            self.assertEqual(cell.border.bottom.style, "thin")
            self.assertEqual(cell.border.bottom.color.rgb[-6:], "1F3864")

    def test_no_cambria_and_header_not_blue(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        self.assertNotEqual(cell.font.name, "Cambria")
                        self.assertNotEqual(cell.fill.fgColor.rgb, "004472C4")

    def test_only_front_tabs_have_navy_tab_colour(self):
        front = {"Priority Actions", "Findings", "SAP vs SoftwareONE", "Risk Heatmap"}
        for ws in self.wb.worksheets:
            colour = ws.sheet_properties.tabColor
            if ws.title in front:
                self.assertTrue(colour is not None and colour.rgb.endswith("1F3864"), ws.title)
            else:
                self.assertIsNone(colour, ws.title)

    def test_print_setup_a4_landscape_one_page_wide(self):
        ws = self.wb["Findings"]
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(int(ws.page_setup.paperSize), 9)
        self.assertEqual(ws.page_setup.fitToWidth, 1)
        self.assertEqual(ws.print_title_rows, "$1:$1")


class PriorityActionsTests(unittest.TestCase):
    def _sheet(self, result):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "out.xlsx"
        generate(result, path)
        return load_workbook(path)["Priority Actions"]

    def _texts(self, ws):
        return [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]

    @staticmethod
    def _link(cell):
        return (cell.hyperlink.target or cell.hyperlink.location) if cell.hyperlink else None

    def test_health_line_in_b2_and_no_written_summary(self):
        result = build_result()
        result.executive_summary = "The system needs attention soon."
        ws = self._sheet(result)
        self.assertEqual(ws["B2"].value, "Overall system health: ● CRITICAL")
        self.assertNotIn("The system needs attention soon.", self._texts(ws))

    def test_top_actions_numbered_once(self):
        result = build_result()
        result.top_5_priority_actions = ["1. Patch the kernel.", "2) Tune the statement."]
        texts = self._texts(self._sheet(result))
        self.assertIn("1. Patch the kernel.", texts)
        self.assertIn("2. Tune the statement.", texts)
        self.assertIn("Top 5 priority actions", texts)

    def test_action_naming_a_finding_shows_a_linked_ref_in_column_h(self):
        result = build_result()  # F-002 is Critical, so row 2 of Findings
        result.top_5_priority_actions = ["Tune the statement (F-002)."]
        ws = self._sheet(result)
        self.assertEqual(ws["H5"].value, "F-002")
        self.assertEqual(self._link(ws["H5"]), "#'Findings'!A2")
        self.assertEqual(ws["H5"].font.u, "single")

    def test_two_refs_fill_two_rows_next_to_one_merged_text(self):
        result = build_result()
        result.top_5_priority_actions = ["Fix F-001 and F-002 together.", "Second action."]
        ws = self._sheet(result)
        self.assertIn("B5:G6", [str(m) for m in ws.merged_cells.ranges])
        self.assertEqual([ws["H5"].value, ws["H6"].value], ["F-001", "F-002"])
        self.assertEqual(self._link(ws["H5"]), "#'Findings'!A3")
        self.assertEqual(self._link(ws["H6"]), "#'Findings'!A2")
        self.assertEqual(ws["B7"].value, "2. Second action.")  # next action starts below both refs

    def test_unknown_ref_gets_no_cell_and_no_link(self):
        result = build_result()
        result.top_5_priority_actions = ["Look at F-999 first."]
        ws = self._sheet(result)
        self.assertIsNone(ws["H5"].value)
        self.assertIsNone(ws["H5"].hyperlink)

    def test_no_actions_shows_one_grey_line(self):
        result = build_result()
        result.top_5_priority_actions = []
        ws = self._sheet(result)
        self.assertEqual(ws["B5"].value, "No priority actions were produced.")

    def test_no_fills_below_the_title_band(self):
        result = build_result()
        result.top_5_priority_actions = ["Fix F-001 and F-002 together.", "Second action."]
        ws = self._sheet(result)
        for row in ws.iter_rows():
            for cell in row:
                if cell.row == 1:  # the navy title band
                    continue
                self.assertIn(cell.fill.fill_type, (None, ""), cell.coordinate)

    def test_title_has_the_system_id_when_known(self):
        result = build_result()
        result.system_profile = SystemProfile(sid="PRD", product="SAP S/4HANA")
        self.assertEqual(self._sheet(result)["B1"].value, "SAP EarlyWatch Alert — PRD")

    def test_title_falls_back_when_the_system_id_is_unknown(self):
        expected = "SAP EarlyWatch Alert — Deep Analysis Report"
        self.assertEqual(self._sheet(build_result())["B1"].value, expected)
        result = build_result()
        result.system_profile = SystemProfile(product="SAP S/4HANA")
        self.assertEqual(self._sheet(result)["B1"].value, expected)

    def _risk_result(self):
        result = build_result()
        result.top_5_priority_actions = ["Tune the statement (F-002)."]
        result.cross_references = [CrossReference(
            title="Old kernel and slow SQL", related_findings=["F-001", "F-999", "F-002"],
            correlation_description="Both hurt response time.", combined_impact="Slow system.",
            recommended_action="Fix both together.")]
        return result

    def test_risks_that_add_up_block_follows_the_top_actions(self):
        ws = self._sheet(self._risk_result())
        self.assertEqual(ws["B6"].value, None)  # one blank row after the single action
        self.assertEqual(ws["B7"].value, "Risks that add up")
        self.assertEqual(ws["B8"].value, "Old kernel and slow SQL\n"
                                         "Why it matters together: Both hurt response time. Slow system.\n"
                                         "What to do: Fix both together.")
        self.assertEqual(ws["B7"].font.b, True)

    def test_risk_refs_are_linked_and_unknown_ids_are_skipped(self):
        ws = self._sheet(self._risk_result())
        self.assertIn("B8:G9", [str(m) for m in ws.merged_cells.ranges])
        self.assertEqual([ws["H8"].value, ws["H9"].value], ["F-001", "F-002"])
        self.assertEqual(self._link(ws["H8"]), "#'Findings'!A3")
        self.assertEqual(self._link(ws["H9"]), "#'Findings'!A2")
        self.assertIsNone(ws["H10"].value)

    def test_risk_title_is_bold_in_rich_text(self):
        from openpyxl.cell.rich_text import CellRichText
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            generate(self._risk_result(), path)
            ws = load_workbook(path, rich_text=True)["Priority Actions"]
        value = ws["B8"].value
        self.assertIsInstance(value, CellRichText)
        parts = [p for p in value if not isinstance(p, str)]
        self.assertTrue(parts[0].font.b)
        self.assertEqual(parts[0].text, "Old kernel and slow SQL")
        self.assertTrue(any(p.text == "What to do: " and p.font.b for p in parts))

    def test_no_cross_references_means_no_block(self):
        texts = self._texts(self._sheet(build_result()))
        self.assertNotIn("Risks that add up", texts)

    def test_risk_without_known_refs_still_shows_one_row(self):
        result = self._risk_result()
        result.cross_references[0].related_findings = ["F-999"]
        ws = self._sheet(result)
        self.assertIn("Old kernel and slow SQL", ws["B8"].value)
        self.assertIsNone(ws["H8"].value)

    def test_there_is_no_linked_issues_tab(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            generate(self._risk_result(), path)
            self.assertNotIn("Linked Issues", load_workbook(path).sheetnames)

    def _coverage_result(self):
        result = build_result()  # sections "1" and "2" have findings
        result.coverage = [
            SectionCoverage(section_id="1", section_title="Software Configuration", status="analyzed"),
            SectionCoverage(section_id="2", section_title="Performance Indicators", status="analyzed"),
            SectionCoverage(section_id="3", section_title="Hardware", status="analyzed"),
            SectionCoverage(section_id="4", section_title="Database", status="analyzed"),
            SectionCoverage(section_id="5", section_title="Security", status="failed"),
        ]
        return result

    def test_checked_line_counts_clean_sections_and_sap_flags(self):
        result = self._coverage_result()
        result.reconciliation = [
            _row("Check", "Not covered", sap_topic="Database", sap_rating="YELLOW", our_section_id="4"),
            _row("Check", "Agree", sap_topic="Hardware", sap_rating="GREEN", our_section_id="3"),
        ]
        texts = self._texts(self._sheet(result))
        self.assertIn("We checked 4 sections, no issues in 2 of them "
                      "(SAP rated 1 red/yellow; see SAP vs SoftwareONE).", texts)

    def test_checked_line_drops_the_sap_part_without_sap_flags(self):
        texts = self._texts(self._sheet(self._coverage_result()))
        self.assertIn("We checked 4 sections, no issues in 2 of them.", texts)

    def test_failed_sections_get_a_bold_line(self):
        ws = self._sheet(self._coverage_result())
        cell = next(c for row in ws.iter_rows() for c in row
                    if c.value == "1 section could not be analysed.")
        self.assertTrue(cell.font.b)
        result = self._coverage_result()
        result.coverage[0].status = "failed"
        self.assertIn("2 sections could not be analysed.", self._texts(self._sheet(result)))

    def test_no_coverage_means_no_checked_or_failed_line(self):
        texts = " ".join(self._texts(self._sheet(build_result())))
        self.assertNotIn("We checked", texts)
        self.assertNotIn("could not be analysed", texts)

    def test_all_sections_failed_shows_only_the_failed_line(self):
        result = build_result()
        result.coverage = [SectionCoverage(section_id="5", section_title="Security", status="failed")]
        texts = self._texts(self._sheet(result))
        self.assertIn("1 section could not be analysed.", texts)
        self.assertFalse(any(t.startswith("We checked") for t in texts))

    def test_how_to_read_block_sits_below_the_other_content(self):
        ws = self._sheet(build_result())
        texts = self._texts(ws)
        head = next(c for row in ws.iter_rows() for c in row if c.value == "How to read this workbook")
        self.assertTrue(head.font.b)
        self.assertGreater(head.row, next(c.row for row in ws.iter_rows() for c in row
                                          if c.value == "1. Tune the statement."))
        legend = " ".join(texts)
        for expected in ("● Critical", "● Low", "●●● High", "●●○ Medium", "●○○ Low"):
            self.assertIn(expected, legend)
        findings_line = next(t for t in texts if t.startswith("Findings —"))
        self.assertIn("Folded columns", findings_line)

    def test_how_to_read_has_a_line_only_for_tabs_that_exist(self):
        texts = " ".join(self._texts(self._sheet(build_result())))
        self.assertNotIn("SAP vs SoftwareONE —", texts)  # no reconciliation rows in this result
        self.assertIn("Risk Heatmap —", texts)
        result = build_result()
        result.reconciliation = [_row("Check", "Agree", sap_rating="GREEN")]
        self.assertIn("SAP vs SoftwareONE —", " ".join(self._texts(self._sheet(result))))
        for da in result.domain_analyses:
            da.findings = []
        self.assertNotIn("Risk Heatmap —", " ".join(self._texts(self._sheet(result))))


if __name__ == "__main__":
    unittest.main()


class RiskHeatmapTests(unittest.TestCase):
    def _generate(self, result):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "out.xlsx"
        generate(result, path)
        return load_workbook(path)

    @staticmethod
    def _find_row(ws, label):
        return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == label)

    def _many_sections(self, n=26):
        result = build_result()
        result.domain_analyses = [
            DomainAnalysis(
                section_title=f"{(i // 5) + 1}.{i % 5 + 1} Part {i}", section_id=f"s{i}",
                findings=[Finding(
                    id=f"G-{i}", title="t", severity="High" if i % 2 else "Low", description="d",
                    evidence="e", impact="i", remediation=_remediation("a"),
                    risk_types=["Security", "Cost"])],
                overall_health="Warning")
            for i in range(n)]
        return result

    def test_totals_match_findings_and_types(self):
        result = build_result()  # F-001 untagged (Medium); F-002 Performance + Cost (Critical)
        ws = self._generate(result)["Risk Heatmap"]
        headers = [c.value for c in ws[3]]
        self.assertEqual(headers[0], "Section")
        self.assertNotIn("SAP rating", headers)  # old payload: no reconciliation
        total = ws[self._find_row(ws, "Total")]
        by_name = {h: c.value for h, c in zip(headers, total)}
        all_rows = self._generate(result)["Findings"].max_row - 1
        self.assertEqual(by_name["Findings"], all_rows)
        self.assertEqual(by_name["Performance"], 1)
        self.assertEqual(by_name["Cost"], 1)
        self.assertEqual(by_name["Not tagged"], 1)
        self.assertIsNone(by_name["Security"])

    def test_cell_fill_is_worst_severity_colour(self):
        from ewa_pipeline.report import excel_style as xs
        result = build_result()
        result.domain_analyses[1].findings.append(Finding(
            id="F-003", title="t", severity="Low", description="d", evidence="e", impact="i",
            remediation=_remediation("a"), risk_types=["Cost"]))
        ws = self._generate(result)["Risk Heatmap"]
        headers = [c.value for c in ws[3]]
        row = self._find_row(ws, "Performance Indicators")
        cost = ws.cell(row, headers.index("Cost") + 1)
        self.assertEqual(cost.value, 2)
        self.assertEqual(cost.fill.start_color.rgb[-6:], xs.SEVERITY_COLOURS["Critical"])
        security = ws.cell(row, headers.index("Security") + 1)
        self.assertIsNone(security.value)
        self.assertIn(security.fill.fill_type, (None, ""))

    def test_sap_red_section_without_findings_appears(self):
        result = build_result()
        result.coverage = [SectionCoverage(section_id="4", section_title="Database", status="analyzed")]
        result.reconciliation = [_row("Check", "Not covered", sap_topic="Database", sap_rating="RED",
                                      our_section_id="4")]
        ws = self._generate(result)["Risk Heatmap"]
        headers = [c.value for c in ws[3]]
        self.assertEqual(headers[1], "SAP rating")
        row = self._find_row(ws, "Database")
        self.assertEqual(ws.cell(row, 2).value, "● Red")
        self.assertEqual(ws.cell(row, headers.index("Findings") + 1).value, "—")

    def test_many_sections_are_grouped_by_chapter(self):
        wb = self._generate(self._many_sections(26))
        ws = wb["Risk Heatmap"]
        self.assertFalse(ws.sheet_properties.outlinePr.summaryBelow)
        hidden = [r for r in range(4, ws.max_row + 1)
                  if ws.row_dimensions[r].outlineLevel == 1 and ws.row_dimensions[r].hidden]
        self.assertEqual(len(hidden), 26)
        chapters = [r for r in range(4, ws.max_row + 1)
                    if ws.row_dimensions[r].outlineLevel == 0 and ws.cell(r, 1).font.bold
                    and ws.cell(r, 1).value != "Total"]
        self.assertEqual(len(chapters), 6)  # 26 sections, 5 per chapter
        total = ws[self._find_row(ws, "Total")]
        headers = [c.value for c in ws[3]]
        self.assertEqual(total[headers.index("Findings")].value, 26)
        self.assertEqual(total[headers.index("Security")].value, 26)

    def test_section_row_links_to_its_first_finding_on_the_findings_tab(self):
        ws = self._generate(build_result())["Risk Heatmap"]
        link = ws.cell(self._find_row(ws, "Performance Indicators"), 1).hyperlink  # Critical: Findings row 2
        self.assertEqual(link.target or link.location, "#'Findings'!A2")
        link = ws.cell(self._find_row(ws, "Software Configuration"), 1).hyperlink  # Medium: Findings row 3
        self.assertEqual(link.target or link.location, "#'Findings'!A3")

    def test_chapter_rows_and_sections_without_findings_have_no_link(self):
        result = self._many_sections(26)
        result.coverage = [SectionCoverage(section_id="4", section_title="Database", status="analyzed")]
        result.reconciliation = [_row("Check", "Not covered", sap_topic="Database", sap_rating="RED",
                                      our_section_id="4")]
        ws = self._generate(result)["Risk Heatmap"]
        chapters = [r for r in range(4, ws.max_row + 1)
                    if ws.row_dimensions[r].outlineLevel == 0 and ws.cell(r, 1).font.bold]
        self.assertTrue(chapters)
        for r in chapters:
            self.assertIsNone(ws.cell(r, 1).hyperlink, r)
        db = self._find_row(ws, "Database")
        self.assertIsNone(ws.cell(db, 1).hyperlink)
        linked = self._find_row(ws, "1.2 Part 1")
        self.assertIsNotNone(ws.cell(linked, 1).hyperlink)

    def test_few_sections_are_not_grouped(self):
        ws = self._generate(self._many_sections(5))["Risk Heatmap"]
        self.assertTrue(all(ws.row_dimensions[r].outlineLevel == 0 for r in range(4, ws.max_row + 1)))

    def test_no_findings_and_no_sap_flags_means_no_tab(self):
        result = build_result()
        for da in result.domain_analyses:
            da.findings = []
        self.assertNotIn("Risk Heatmap", self._generate(result).sheetnames)

    def test_tab_position_after_sap_vs_softwareone(self):
        result = build_result()
        result.coverage = [SectionCoverage(section_id="9", section_title="Clean", status="analyzed")]
        result.reconciliation = [_row("Check", "Agree", sap_rating="GREEN")]
        names = self._generate(result).sheetnames
        self.assertEqual(names[1:4], ["Findings", "SAP vs SoftwareONE", "Risk Heatmap"])
